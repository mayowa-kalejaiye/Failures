"""Front matter + Chapter 1 of the Failures final-year project report."""
import sys

sys.path.insert(0, str(__import__("pathlib").Path(__file__).resolve().parent))
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from common import (CITED, OUT, STAGE_DIR, body, chapter_head, cite, front_head, new_document, page_break,
                    para, section_head)

TITLE = ("DESIGN AND IMPLEMENTATION OF A DETERMINISTIC FAILURE-MODE "
         "GUARDRAIL SYSTEM FOR AI-GENERATED SOFTWARE USING MODEL "
         "CONTEXT PROTOCOL")

front = new_document()
doc = new_document()

# ---------------------------------------------------------------- title page
para(front, TITLE, bold=True, center=True, size=14, space_before=72,
     space_after=24)
para(front, "BY:", bold=True, center=True, space_after=6)
para(front, "OLUWAMAYOWA KALEJAIYE", bold=True, center=True,
     space_after=2)
para(front, "2023/C/SENG/0351", bold=True, center=True, space_after=24)
para(front, "DEPARTMENT OF SOFTWARE ENGINEERING,", bold=True, center=True,
     space_after=2)
para(front, "SCHOOL OF COMPUTING,", bold=True, center=True, space_after=2)
para(front, "MIVA OPEN UNIVERSITY ABUJA,", bold=True, center=True,
     space_after=2)
para(front, "NIGERIA", bold=True, center=True, space_after=36)
para(front, "IN PARTIAL FULFILMENT OF THE REQUIREMENTS FOR THE AWARD OF THE",
     bold=True, center=True, space_after=2)
para(front, "BACHELOR OF SCIENCE (HONOURS) DEGREE IN SOFTWARE ENGINEERING",
     bold=True, center=True, space_after=36)
para(front, "SUPERVISED BY: IJEGWA ACHEME Ph.D", bold=True, center=True,
     space_after=36)
para(front, "JULY, 2026", bold=True, center=True)
page_break(doc)

# ------------------------------------------------------------ certification
front_head(front, "CERTIFICATION")
body(front, "This is to certify that I am responsible for the work submitted in "
          "this project, that the original work is mine, except as specified in "
          "the acknowledgement and references. And that neither the project nor "
          "the original work contained therein has been submitted to the "
          "University or any other institution for the award of a degree.")
para(front, "Oluwamayowa Kalejaiye Samuel, 2023/C/SENG/0351", center=True,
     space_before=24, space_after=0, justify=False)
para(front, "STUDENT\u2019S NAME AND MATRIC NUMBER", center=True, justify=False)
page_break(doc)

# ---------------------------------------------------------------- approval
front_head(front, "APPROVAL")
body(front, "This project has been approved for the Department of Software "
          "Engineering, School of Computing, Miva Open University, Abuja, Nigeria.")
for name, role in [
        ("Ijegwa Acheme, Ph.D", "Name of Supervisor"),
        ("Morolake O. Lawrence, Ph.D, MHCPN, MNSC", "Head of Department"),
        ("Emeka Emmanuel Ogbuju, Ph.D", "Dean, School of Computing"),
        ("Name of External Examiner", "External Examiner")]:
    para(front, name, bold=True, space_before=18, space_after=0, justify=False)
    para(front, role, space_after=0, justify=False)
    para(front, "Signature and Date ......................................",
         space_after=6, justify=False)
page_break(doc)

# --------------------------------------------------------------- dedication
front_head(front, "DEDICATION")
body(front, "This project is dedicated to God Almighty for the abundant grace, "
          "wisdom, knowledge, and skills given to me all through my life, "
          "especially during my stay at Miva Open University, Abuja.")
page_break(doc)

# ---------------------------------------------------------- acknowledgement
front_head(front, "ACKNOWLEDGEMENT")
body(front, "I would like to express my profound gratitude to God Almighty for "
          "His grace throughout this research. My sincere appreciation goes to "
          "my supervisor, Ijegwa Acheme, for his invaluable guidance, "
          "continuous support, and constructive feedback which contributed "
          "immensely to the success of this project.")
body(front, "I also want to thank the Head of Department, the Dean of the School "
          "of Computing, and all the academic and non-academic staff of the "
          "Department of Software Engineering, Miva Open University, for their "
          "efforts in ensuring a conducive learning environment.")
page_break(doc)

# ================================================================ CHAPTER 1
front.save(str(STAGE_DIR / "stage_front.docx"))
print("Saved stage_front")
chapter_head(doc, "CHAPTER ONE: INTRODUCTION")

section_head(doc, "1.1 Background of the Study")
body(doc, "Software development has entered a period in which large language "
          "models trained on source code contribute directly to production "
          "systems. Chen et al. (2021) established functional-correctness "
          "evaluation of code-trained models, demonstrating that models could "
          "solve a substantial share of self-contained programming problems; "
          "subsequent deployment studies measured productivity gains from "
          "GitHub Copilot in professional practice (Peng, Kalliamvakou, Cihon, "
          "& Demirer, 2023). Coding agents extend this trajectory from "
          "suggesting code to scaffolding complete systems: given a natural "
          "language prompt, an agent now produces payment handlers, "
          "authentication flows, and background workers that run. The question "
          "facing the field is therefore no longer whether generated code runs, "
          "but whether the systems it composes survive contact with "
          "production.")
cite(doc, "chen2021", "peng2023")
body(doc, "The evaluation culture that grew around these models measures "
          "functional correctness almost exclusively: pass rates on held-out "
          "programming problems (Chen et al., 2021) and resolution rates on "
          "real repository issues (Jimenez et al., 2024). In parallel, agents "
          "have been extended through tool use, and the Model Context Protocol "
          "has standardised how a model discovers and invokes external tools "
          "(Anthropic, 2024). The combination is powerful: an agent with tool "
          "access can design, implement, and test a feature end to end. But "
          "the checks available at each of those steps verify behaviour on the "
          "paths that were exercised. Nothing in the standard workflow asks "
          "what happens on the paths that were not.")
cite(doc, "jimenez2024", "anthropic2024")
body(doc, "The paths that were not exercised are where production incidents "
          "live. The most dangerous of them is the ambiguous outcome: the "
          "operation may or may not have taken effect, and the caller cannot "
          "tell which. A payment provider that processes a charge and then "
          "loses the response to a network partition leaves the caller "
          "unable to distinguish success from failure, because a timeout "
          "reveals nothing about what happened on the other side of the "
          "network (Kleppmann, 2017). The agent-written handler, like most "
          "hand-written tutorials it learned from, treats that ambiguity as "
          "impossible: it either retries blindly and double-charges the "
          "customer, or it abandons the payment and loses the enrollment. "
          "Neither branch raises an error. Stability engineering treats "
          "exactly these cases as first-class design concerns (Nygard, 2018), "
          "and the site reliability discipline exists largely to learn from "
          "them systematically through postmortems (Beyer, Jones, Petoff, & "
          "Murphy, 2016).")
cite(doc, "kleppmann2017", "nygard2018", "beyer2016")
body(doc, "Previous research documents each side of this problem without "
          "closing it. Benchmarks continue to measure correctness rather "
          "than resilience (Chen et al., 2021; Jimenez et al., 2024). Security "
          "audits of generated code find a substantial share of suggestions "
          "vulnerable (Pearce, Ahmad, Tan, Dolan-Gavitt, & Karri, 2022). "
          "Systematic reviews of language models in architecture tasks report "
          "hallucination as a central limitation (Schmid, Verbeek, & Van Der "
          "Storm, 2025). The resilience literature catalogues the controls "
          "that address these failures \u2014 idempotency, outboxes, "
          "reconciliation, bulkheads \u2014 but as engineering knowledge, not "
          "as machine-checkable properties (Nygard, 2018; Kleppmann, 2017). "
          "Decision support systems compare technologies once criteria are "
          "defined (Farshidi & Jansen, 2020), and studies of architectural "
          "practice document how rarely decisions and their rationale are "
          "recorded at all (Capilla, Jansen, Tang, Avgeriou, & Babar, 2016).")
cite(doc, "pearce2022", "schmid2025", "farshidi2020", "capilla2016")
body(doc, "The problem is not confined to large technology firms. Surveys of "
          "software practitioners in Nigerian universities find that "
          "well-documented architectural requirements materially aid system "
          "construction, while also showing how narrow and local most such "
          "evaluations remain (Bernard, Ayo, & Odunaike, 2020). For student "
          "developers and small teams, who cannot afford a reliability "
          "organisation, the absence of checkable guidance is felt most "
          "acutely: the same ambiguous outcome that a large firm absorbs "
          "through process can end a small product outright. A guardrail "
          "system that is deterministic, free to run, and explainable in its "
          "findings is therefore disproportionately valuable where "
          "engineering mentorship is scarcest.")
cite(doc, "bernard2020")
body(doc, "What does not yet exist is a deterministic layer between the agent "
          "and its output that checks failure modes before code ships, with "
          "evidence, and at zero model cost. Evaluation practice itself shows "
          "why such a layer is missing: a survey of software architecture "
          "research found that most assessments rest on experiments and case "
          "studies, while only a small minority of papers ship replication "
          "packages (Konersmann et al., 2022). Findings that cannot be "
          "reproduced cannot be checked, and checks that require a second "
          "model call inherit the unreliability they are meant to guard "
          "against.")
cite(doc, "konersmann2022")
body(doc, "The measurement culture reinforces the blind spot. Benchmarks that "
          "made code generation a respectable research area all score "
          "functional outcomes: whether the program passes hidden tests, "
          "whether the issue is resolved, whether the suggestion compiles. "
          "Jimenez et al. (2024) moved the field from synthetic problems to "
          "real repository issues, which was genuine progress, but the "
          "scoring axis did not move with it: a resolved issue is still a "
          "functioning issue, not a resilient one. A payment handler that "
          "passes every test and double-charges on retry is a success under "
          "every standard metric in use. The consequence is that agents "
          "optimise, implicitly, for the demonstration rather than the "
          "deployment \u2014 for code that convinces a reviewer in the "
          "moment rather than code that survives the network at 3 a.m.")
cite(doc, "jimenez2024")
body(doc, "The engineering literature, by contrast, has treated failure as "
          "the primary object of study for decades. Nygard (2018) catalogues "
          "stability patterns \u2014 bulkheads, circuit breakers, timeouts "
          "\u2014 each one a response to a specific production failure the "
          "author had witnessed. Kleppmann (2017) builds the theoretical "
          "account: unbounded delay makes failure detection impossible, "
          "exactly-once semantics unattainable in the general case, and "
          "idempotence the practical substitute. Beyer et al. (2016) "
          "institutionalise the learning loop, turning incidents into "
          "postmortems and postmortems into design constraints. What unites "
          "these sources is a premise the benchmarks do not share: that the "
          "interesting behaviour of a system is what it does when something "
          "goes wrong, and that this behaviour must be designed rather than "
          "hoped for.")
cite(doc, "nygard2018", "kleppmann2017", "beyer2016")
body(doc, "This study proposes that layer. Failures encodes eleven failure "
          "dimensions \u2014 atomicity, idempotency, timeout and ambiguous "
          "outcome, concurrency, ordering, consistency, availability, "
          "resource exhaustion, recovery, observability, and retry safety "
          "\u2014 as testable invariants backed by deterministic static "
          "rules, exposes them to coding agents as thirteen Model Context "
          "Protocol tools, and scores the resulting systems with a "
          "deterministic harness across ten resilience criteria. A blind "
          "study across three commercial agents over six production-style "
          "scenarios measures whether the same prompt, issued with and "
          "without the guardrails, produces measurably more resilient "
          "software (Kalejaiye, 2026).")
cite(doc, "kalejaiye2026")

section_head(doc, "1.2 Statement of the Problem")
body(doc, "Although code-generating models have improved the speed at which "
          "software is produced, several challenges remain. These include "
          "silent mishandling of ambiguous outcomes, missing idempotency "
          "controls on side-effecting calls, and absent transaction "
          "boundaries around multi-step operations. Existing approaches such "
          "as functional-correctness benchmarks (Chen et al., 2021; Jimenez "
          "et al., 2024) and general-purpose static analysers have attempted "
          "to assure quality; however, they suffer from complementary "
          "limitations. Benchmarks verify behaviour only on exercised paths, "
          "while analysers encode style and syntax rather than engineering "
          "intent, and language-model judges hallucinate defects at rates "
          "that teach engineers to dismiss warnings (Pearce et al., 2022; "
          "Schmid et al., 2025). As a result, developers and organisations "
          "continue to experience production incidents \u2014 duplicate "
          "charges, lost updates, and unrecoverable partial states \u2014 "
          "that no exercised test foresaw (Nygard, 2018; Kleppmann, 2017). "
          "Therefore, there is a need to develop a deterministic "
          "failure-mode guardrail system, delivered to coding agents as "
          "tools, capable of checking resilience properties with line-level "
          "evidence before code ships.")
cite(doc, "chen2021", "jimenez2024", "pearce2022", "schmid2025",
     "nygard2018", "kleppmann2017")

section_head(doc, "1.3 Aim and Objectives of the Study")
para(doc, "Aim of the Study", bold=True, justify=False)
body(doc, "The primary aim of this study is to design, implement, and "
          "evaluate a deterministic failure-mode guardrail system that "
          "improves the failure resilience of AI-generated software.")
para(doc, "Objectives of the Study", bold=True, justify=False)
body(doc, "The objectives of this study are to:")
for o in [
    "To investigate failure modes in AI-generated software and existing "
    "guardrail approaches.",
    "To design a deterministic failure-mode guardrail framework based on "
    "engineering invariants.",
    "To develop a Model Context Protocol server implementing the guardrail "
    "rules with evidence-backed findings.",
    "To implement a deterministic evaluation harness for scoring the "
    "resilience of agent-generated software.",
    "To evaluate the guardrail system through proxy and blind agent studies "
    "and analyse the results."]:
    para(doc, o, justify=False)
para(doc, "Research Questions", bold=True, justify=False)
for q in [
    "What failure modes occur in AI-generated software, and how do existing "
    "approaches address them?",
    "How can engineering invariants be formalised as deterministic, "
    "checkable guardrails?",
    "How can the guardrail framework be implemented as a Model Context "
    "Protocol server that returns evidence-backed findings?",
    "How can a deterministic harness score the resilience of "
    "agent-generated software across scenarios?",
    "To what extent does connecting the guardrail system improve the "
    "failure resilience of agent-generated software?"]:
    para(doc, q, justify=False)

section_head(doc, "1.4 Significance of the Study")
body(doc, "The findings of this study will benefit software developers by "
          "providing guardrails that catch resilience defects during "
          "development rather than in production, where ambiguous outcomes "
          "become support tickets instead of pages (Nygard, 2018). The study "
          "will benefit industry practitioners and organisations by reducing "
          "the class of incidents \u2014 duplicate charges, lost updates, "
          "unrecoverable partial states \u2014 that arise when generated code "
          "meets unreliable networks (Kleppmann, 2017). The study will "
          "benefit researchers by contributing a reproducible evaluation "
          "method: frozen prompts, hashed manifests, and a deterministic "
          "harness whose tables regenerate from artifacts rather than from "
          "transcription (Konersmann et al., 2022). The study will benefit "
          "students of software engineering by making failure reasoning "
          "explicit and checkable instead of tacit and experiential.")
cite(doc, "nygard2018", "kleppmann2017", "konersmann2022")

section_head(doc, "1.5 Scope of the Study")
body(doc, "This study focuses on the design, implementation, and evaluation "
          "of a deterministic failure-mode guardrail system delivered to "
          "coding agents through the Model Context Protocol. Specifically, "
          "the study covers eleven failure dimensions, thirteen guardrail "
          "tools, a ten-criterion evaluation harness, six production-style "
          "scenarios, and a blind comparison across three commercial coding "
          "agents. However, the study excludes general-purpose defect "
          "detection unrelated to failure resilience, dynamic testing against "
          "live failing providers, and languages other than Python, because "
          "the evaluation corpus and the rule engine are implemented for a "
          "single language and a fixed scenario set.")

section_head(doc, "1.6 Limitations of the Study")
body(doc, "The study was limited by the heuristic nature of the rule engine, "
          "which reasons over source structure rather than interprocedural "
          "data flow and therefore cannot prove lock scope or acknowledgement "
          "ordering. The study was limited by time constraints, which did not "
          "allow dynamic execution of generated systems against failing "
          "providers. These constraints were managed through honest "
          "confidence labelling on every finding, an adversarial set that "
          "documents exactly where heuristics stop, and a scope restricted to "
          "static guarantees the engine can actually support.")

section_head(doc, "1.7 Definition of Terms")
for term, text in [
    ("Failure: ", "any deviation in which the system cannot guarantee that "
     "an operation\u2019s intended effect occurred exactly once and the "
     "system remains in a valid state. Failures of the ambiguous kind, in "
     "which a timeout reveals nothing about whether a remote effect "
     "occurred, are the central concern of distributed systems design "
     "(Kleppmann, 2017)."),
    ("Ambiguous outcome: ", "a failure mode in which the caller does not "
     "know whether a side effect happened, arising from timeouts, crashes "
     "before acknowledgement, or network partitions; the most dangerous kind "
     "because the system is consistent but its knowledge of that state is "
     "wrong (Kleppmann, 2017; Kalejaiye, 2026)."),
    ("Idempotency: ", "the property that one logical operation produces at "
     "most one effect regardless of how many times it executes; the standard "
     "control is an idempotency key persisted before the side effect "
     "(Kleppmann, 2017)."),
    ("Engineering invariant: ", "a predicate that must hold under any "
     "failure, such as that a payment is charged at most once per logical "
     "operation (Kalejaiye, 2026)."),
    ("Model Context Protocol: ", "an open protocol standardising how a "
     "language model discovers and invokes external tools, used here as the "
     "delivery mechanism for guardrail checks (Anthropic, 2024)."),
    ("Coding agent: ", "a language-model system that writes, modifies, and "
     "tests software through tool use, evaluated in this study on its "
     "ability to produce failure-resilient systems rather than merely "
     "functioning code (Chen et al., 2021; Jimenez et al., 2024)."),
    ("Static failure check: ", "an automated examination of source code "
     "without execution that reports a potential invariant violation with "
     "line-level evidence and a calibrated confidence, never labelled as a "
     "proven defect unless syntactically provable (Kalejaiye, 2026)."),
    ("Resilience: ", "the ability of a system to preserve its invariants "
     "under failures, achieved through explicit controls rather than "
     "optimistic assumptions (Nygard, 2018; Beyer et al., 2016)."),
]:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(term)
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r2 = p.add_run(text)
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(12)
cite(doc, "kleppmann2017", "kalejaiye2026", "anthropic2024", "chen2021",
     "jimenez2024", "nygard2018", "beyer2016")

doc.save(str(STAGE_DIR / "stage_ch1.docx"))
print("Saved stage_ch1")
(STAGE_DIR / "stage_ch1.cited.json").write_text(__import__("json").dumps(CITED), encoding="utf-8")
