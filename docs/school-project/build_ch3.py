"""Chapter 3 of the Failures final-year project report. Stages to stage_ch3."""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from common import (CITED, OUT, STAGE_DIR, body, chapter_head, cite, new_document,
                    page_break, section_head)

STAGE3 = STAGE_DIR / "stage_ch3.docx"
FIG = OUT.parent / "figures"
doc = new_document()
page_break(doc)
chapter_head(doc, "CHAPTER THREE: RESEARCH METHODOLOGY")


def fig(doc, fname, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(FIG / fname), width=Inches(5.9))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(10)
    r = c.add_run(caption)
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)


def table(doc, header, rows, widths=None, size=10, caption=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = "Table Grid"
    for j, h in enumerate(header):
        cell = t.cell(0, j)
        cell.text = ""
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.bold = True
        r.font.name = "Times New Roman"
        r.font.size = Pt(size)
        if widths:
            cell.width = int(widths[j] * 914400)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Times New Roman"
            r.font.size = Pt(size)
            if widths:
                cell.width = int(widths[j] * 914400)
    if caption:
        from docx.enum.text import WD_ALIGN_PARAGRAPH as _A
        from docx.shared import Pt as _Pt
        c = doc.add_paragraph()
        c.alignment = _A.CENTER
        c.paragraph_format.space_after = _Pt(10)
        r = c.add_run(caption)
        r.bold = True
        r.font.name = "Times New Roman"
        r.font.size = _Pt(11)
    doc.add_paragraph()


def code(doc, lines):
    for ln in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.left_indent = Inches(0.4)
        r = p.add_run(ln if ln.strip() else " ")
        r.font.name = "Courier New"
        r.font.size = Pt(9.5)
    doc.add_paragraph()


# ------------------------------------------------------------------ 3.1
section_head(doc, "3.1 Introduction")
body(doc, "This chapter describes the methodology adopted to achieve the "
          "objectives of this study. It presents the research design, "
          "development methodology, system requirements, data collection "
          "methods, system design, implementation tools, testing procedures, "
          "and evaluation techniques used in developing and assessing the "
          "proposed guardrail system.")

# ------------------------------------------------------------------ 3.2
section_head(doc, "3.2 Research Methodology (Research Design)")
body(doc, "This study adopted a Design Science Research approach because the "
          "primary objective was to develop and evaluate a software solution "
          "for the failure resilience of AI-generated code. This approach is "
          "appropriate because it supports system design, implementation, "
          "testing, and performance evaluation within a single coherent "
          "frame: the artefact is built, then the artefact is measured.")
body(doc, "Design Science Research develops and evaluates an innovative "
          "computing artefact rather than testing behavioural hypotheses, "
          "which makes it the standard design for studies whose contribution "
          "is a working system plus evidence that the system works (Hevner, "
          "March, Park, & Ram, 2004; Peffers, Tuunanen, Rothenberger, & "
          "Chatterjee, 2007). Alternative designs were considered and "
          "rejected on fit. Experimental research compares algorithms under "
          "controlled conditions, but this study compares development "
          "conditions (guardrails on or off) rather than algorithms. Survey "
          "research collects opinions, but opinions about resilience are a "
          "poor substitute for measured findings on code. Case study "
          "research investigates one organisation in depth, but a single "
          "site cannot answer whether an intervention generalises across "
          "agents. The study therefore follows Design Science Research: "
          "build the guardrail server as the artefact, then evaluate it "
          "through controlled proxy and blind agent studies whose results "
          "are reproducible from frozen artifacts.")
cite(doc, "hevner2004", "peffers2007")

# ------------------------------------------------------------------ 3.3
section_head(doc, "3.3 System Development Methodology")
body(doc, "The system itself was developed using Agile Development, which "
          "organises work into short iterations of design, implementation, "
          "testing, and review rather than a single sequential pass (Beck et "
          "al., 2001). Each iteration delivered a working slice \u2014 first "
          "the rule engine over a fixed scenario, then the MCP tool surface, "
          "then the harness and its criteria \u2014 so that every component "
          "was exercised against real agent output before the next was "
          "built. Waterfall was rejected because the rule set and the "
          "criteria co-evolved: writing rules revealed which criteria were "
          "measurable, and scoring runs revealed which rules were noisy. "
          "That feedback loop is precisely what iterative development "
          "exists for, and a strictly sequential process would have frozen "
          "both too early.")
cite(doc, "beck2001")
table(doc, ["Methodology", "Suitable For"],
      [["Waterfall Model", "Projects with stable, fully known requirements"],
       ["Agile Development", "Iterative software development (adopted)"],
       ["Scrum", "Agile team-based projects"],
       ["Spiral Model", "High-risk software projects"],
       ["Prototype Model", "User-centred systems requiring frequent feedback"],
       ["DevOps", "Continuous integration and deployment"]],
      widths=[2.2, 4.0],
      caption="Table 3.1: Common System Development Methodologies")

# ------------------------------------------------------------------ 3.4
section_head(doc, "3.4 Requirements Analysis")
section_head(doc, "3.4.1 Functional Requirements")
body(doc, "What the system must do.")
table(doc, ["Requirement ID", "Functional Requirement"],
      [["FR1", "Accept source code or a development plan for review"],
       ["FR2", "Analyse code against ten rule families across failure dimensions"],
       ["FR3", "Return findings with severity, confidence, and line-level evidence"],
       ["FR4", "Expose all checks as typed Model Context Protocol tools"],
       ["FR5", "Check idempotency, retry safety, and transaction safety on demand"],
       ["FR6", "Generate failure tests mapped to each finding"],
       ["FR7", "Score finished systems against ten resilience criteria"],
       ["FR8", "Freeze prompts, manifests, and commits for reproducibility"]],
      widths=[1.2, 5.0],
      caption="Table 3.2: Functional Requirements")
section_head(doc, "3.4.2 Non-Functional Requirements")
table(doc, ["Requirement ID", "Requirement"],
      [["NFR1", "Determinism: identical code always yields identical findings"],
       ["NFR2", "Zero model-token cost: no LLM in the checking loop"],
       ["NFR3", "Honesty: confidence-labelled output, never overstated"],
       ["NFR4", "Reproducibility: every table regenerates from artifacts"],
       ["NFR5", "Portability: one server serves any MCP-capable agent"]],
      widths=[1.2, 5.0],
      caption="Table 3.3: Non-Functional Requirements")

# ------------------------------------------------------------------ 3.5
section_head(doc, "3.5 Data Collection")
body(doc, "The data in this study are experimental data generated by running "
          "coding agents under controlled conditions, not survey responses "
          "or scraped records. What was collected: matched pairs of agent "
          "output (baseline without guardrails, enabled with guardrails) for "
          "six production-style scenarios across three commercial agents, "
          "plus six adversarial programs and six hand-built proxy pairs. Why "
          "these data: only matched pairs isolate the single manipulated "
          "variable, which is whether the guardrail server was connected. "
          "Where obtained: generated in-house by driving Claude, Cursor, and "
          "Codex with frozen scenario prompts, and by authoring proxy and "
          "adversarial cases directly. How collected: the deterministic "
          "harness (run_evaluation.py) scored every output on the same ten "
          "criteria; seventeen real-agent pairs, six proxy pairs, and six "
          "adversarial cases resulted. Who provided the data: the three "
          "agents for outputs, the author for prompts and adversarial cases. "
          "Data quality was ensured by freezing scenario prompts "
          "(prompt_hash), recording model, timestamp, and commit in every "
          "run manifest, and freezing the harness after the first real "
          "result so later scenarios could not be tuned to the scorer.")
table(doc, ["Data Source", "Description", "Purpose", "Size/Records"],
      [["Agent outputs (matched pairs)", "Baseline vs enabled code per scenario",
        "Isolate the intervention effect", "17 pairs, 3 agents"],
       ["Proxy corpus", "Hand-built naive vs improved implementations",
        "Harness regression test", "6 pairs"],
       ["Adversarial set", "Plausible-looking but unsafe programs",
        "Probe checker limits", "6 cases"],
       ["Run manifests", "Prompt hash, model, timestamp, commit",
        "Reproducibility proof", "17 manifests"]],
      widths=[1.6, 1.8, 1.5, 1.3], size=9,
      caption="Table 3.4: Summary of Data Sources")

# ------------------------------------------------------------------ 3.6
section_head(doc, "3.6 System Analysis")
body(doc, "The existing system is the default agent workflow: prompt in, "
          "code out, judged by functional benchmarks (Chen et al., 2021; "
          "Jimenez et al., 2024), optionally assisted by linters that check "
          "style and by model judges that hallucinate (Schmid et al., 2025). "
          "Its problems: no failure dimension is named, no invariant is "
          "checked, and no evidence attaches to any warning. The proposed "
          "system inserts a deterministic checking layer between the agent "
          "and its output: the same prompt, the same agent, but every "
          "result passes through named, evidenced, confidence-labelled "
          "rules before it counts. Its advantages: prevention instead of "
          "postmortems, zero token cost, and findings a reviewer can verify "
          "line by line.")
cite(doc, "chen2021", "jimenez2024", "schmid2025")
table(doc, ["Aspect", "Existing Workflow", "Proposed System"],
      [["Failure handling", "Unnamed, untested", "11 named dimensions, 10 rules"],
       ["Judgement", "Model-based, stochastic", "Deterministic, reproducible"],
       ["Evidence", "None attached", "Line excerpt + lines on every finding"],
       ["Cost per check", "Model tokens", "Zero tokens"],
       ["Trust model", "Warnings routinely dismissed",
        "Confidence-labelled, never overstated"]],
      widths=[1.4, 2.4, 2.4],
      caption="Table 3.5: Existing Workflow versus Proposed System")

# ------------------------------------------------------------------ 3.7
section_head(doc, "3.7 System Design")
body(doc, "Figure 3.1 shows the system architecture. The coding agent is the "
          "client; the MCP server exposes thirteen typed tools; the "
          "deterministic engine applies ten rule families drawn from the "
          "knowledge base; findings return with severity, confidence, and "
          "evidence; and the evaluation harness scores finished systems "
          "against ten criteria from frozen artifacts.")
fig(doc, "fig31_architecture.png",
    "Figure 3.1: System architecture of Failures.")
body(doc, "Figure 3.2 shows the knowledge graph that drives every check: a "
          "principle names the failure modes it covers, each failure mode "
          "is mitigated by a pattern, and each pattern is verified by a "
          "test. Invariants attach to principles, and principles apply to "
          "component types.")
fig(doc, "fig32_knowledge_graph.png",
    "Figure 3.2: Knowledge graph — Principle → FailureMode → Pattern → Test.")
body(doc, "Figure 3.3 shows the use cases. The coding agent submits code or "
          "plans, receives findings with evidence, generates failure tests, "
          "and views resilience scores. The administrator maintains the "
          "knowledge-base JSON and freezes prompts with manifests.")
fig(doc, "fig33_usecase.png", "Figure 3.3: Use case diagram.")
body(doc, "Figure 3.4 shows the evaluation pipeline: seventeen matched "
          "baseline/enabled pairs, produced from the same six prompts, "
          "scored by one deterministic harness into the results tables of "
          "Chapter Four.")
fig(doc, "fig34_eval_flow.png", "Figure 3.4: Evaluation pipeline.")

# ------------------------------------------------------------------ 3.8
section_head(doc, "3.8 Algorithm / Proposed Model")
body(doc, "The review algorithm scans source once per rule family, attaches "
          "calibrated confidence, and sorts by severity. The harness then "
          "scores each finished system on ten criteria and differences the "
          "matched arms.")
code(doc, [
    "function review_code(source):",
    "    findings = []",
    "    for rule in RULES:            # 10 rule families",
    "        for match in rule.scan(source):",
    "            findings.append({",
    "              id, dimension,",
    "              severity: rule.max_severity,",
    "              confidence: calibrate(match),  # 0.95+ proof, else heuristic",
    "              evidence: {excerpt, lines, note}})",
    "    return sort_by_severity(findings)",
    "",
    "function evaluate(scenario):",
    "    for arm in [baseline, failures_enabled]:",
    "        code = load(scenario, arm)      # frozen prompt + manifest",
    "        score[arm] = ten_criteria(",
    "            review_code(code), check_idempotency(code),",
    "            check_transaction_safety(code),",
    "            check_retry_safety(code))",
    "    return delta(score[enabled], score[baseline])",
])

# ------------------------------------------------------------------ 3.9
section_head(doc, "3.9 Implementation Tools")
table(doc, ["Component", "Tool / Library", "Version", "Purpose"],
      [["Language", "Python", "3.11", "Engine, harness, tooling"],
       ["Agent protocol", "mcp", ">= 2.0", "MCP server surface"],
       ["Validation", "pydantic", ">= 2", "Tool schemas, findings"],
       ["Serving", "anyio / starlette / httpx", "pinned", "Async I/O, transport"],
       ["Testing", "pytest", "8.x", "Harness and agent-suite runs"],
       ["Version control", "Git / GitHub", "—", "Code and manifest history"],
       ["Publishing", "Zenodo", "—", "DOI artifact archive"],
       ["Manuscript", "ReportLab / MiKTeX", "—", "Regenerable paper builds"]],
      widths=[1.3, 1.7, 0.8, 2.4], size=9,
      caption="Table 3.6: Implementation Tools")

# ------------------------------------------------------------------ 3.10
section_head(doc, "3.10 Testing and Evaluation")
body(doc, "Unit testing covered each rule family against positive and "
          "negative snippets. Integration testing ran the full harness end "
          "to end in proxy mode (six scenarios) and manual mode (seventeen "
          "agent pairs). System testing scored complete agent outputs rather "
          "than fragments. User acceptance testing took the form of the "
          "blind agent study itself: three commercial agents used the tools "
          "in normal workflows. Performance testing confirmed checks run in "
          "well under a second with zero model tokens.")
table(doc, ["Test Case", "Expected Result", "Actual Result", "Status"],
      [["TC1 payment proxy", "CRITICAL 3→0", "CRITICAL 3→0", "Pass"],
       ["TC2 webhook proxy", "CRITICAL 3→0", "CRITICAL 3→0", "Pass"],
       ["TC3 queue HIGH", "HIGH 2→0", "HIGH 2→0", "Pass"],
       ["TC4 adversarial idempotency", "Flagged CRITICAL",
        "Flagged CRITICAL", "Pass"],
       ["TC5 adversarial ack order", "Flagged (HIGH)", "Flagged (HIGH)", "Pass"],
       ["TC6 upsert false positive", "Documented, not hidden",
        "Reported in §4/§5", "Pass"]],
      widths=[1.6, 1.5, 1.5, 0.8], size=9,
      caption="Table 3.7: Test Cases and Results")

# ------------------------------------------------------------------ 3.11
section_head(doc, "3.11 Ethical Considerations")
body(doc, "No human participants were involved and no personal data was "
          "collected: all scenarios are synthetic, all outputs are generated "
          "code, and the only credentials in the repository are redacted "
          "test placeholders. The tool itself is advisory by design \u2014 "
          "it reports what its heuristics support, labels the rest with "
          "confidence, and never certifies software as correct \u2014 so "
          "that no user mistakes a check for a guarantee. The system, the "
          "evaluation data, and the paper are released under open licences "
          "(MIT for code, CC BY 4.0 for the manuscript), and the work "
          "proceeded under the supervisor\u2019s oversight with frozen "
          "artifacts that any examiner can re-run.")

section_head(doc, "3.12 Chapter Summary")
body(doc, "This chapter set out the Design Science Research design, the "
          "Agile development process, the functional and non-functional "
          "requirements, the experimental data and its quality controls, the "
          "existing and proposed systems, the architecture with four "
          "figures, the review and scoring algorithms, the implementation "
          "tools, the testing evidence, and the ethical position. The next "
          "chapter presents what the built system produced.")

doc.save(str(STAGE3))
print(f"Saved {STAGE3}")

(STAGE_DIR / "stage_ch3.cited.json").write_text(__import__("json").dumps(CITED), encoding="utf-8")
